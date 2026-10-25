"""Verify synthetic QA exports written by `npm test` (pypdf, pdfplumber, openpyxl)."""

from __future__ import annotations

import json
import re
from pathlib import Path

import pdfplumber
from openpyxl import load_workbook
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
QA = ROOT / ".qa"
SOURCE = json.loads((ROOT / "inputs" / "workbook-extract.json").read_text(encoding="utf-8"))
PBC = SOURCE["sheets"][0]["cells"]
META = SOURCE["sheets"][1]["cells"]


def normalized(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def verify(name: str) -> None:
    pdf_path = QA / f"{name}.pdf"
    xlsx_path = QA / f"{name}.xlsx"
    assert pdf_path.exists() and xlsx_path.exists(), f"Missing {name} QA exports; run npm test first"
    reader = PdfReader(pdf_path)
    pdf_text = normalized(" ".join(page.extract_text() or "" for page in reader.pages))
    workbook = load_workbook(xlsx_path, read_only=True, data_only=False)
    assert len(reader.pages) >= 1
    assert {"Summary", "Check ledger", "PBC coverage", "Assurance criteria", "Evidence", "Metadata", "Exceptions", "Reviewer signoff"}.issubset(workbook.sheetnames)
    pack_id = workbook["Summary"]["B3"].value
    snapshot = workbook["Summary"]["B4"].value
    assert pack_id in pdf_text and snapshot in pdf_text
    assert SOURCE["sha256"] in pdf_text
    assert "Reviewer signoff" in pdf_text
    for item in PBC:
        if item["cell"].startswith("A") and item["cell"] != "A1":
            assert item["value"] in pdf_text, item
    for item in PBC:
        if item["cell"].startswith("F") and item["cell"] != "F1":
            for criterion in item["value"].split(";"):
                assert normalized(criterion) in pdf_text, criterion
    for item in META:
        if item["cell"].startswith("A") and item["cell"] != "A1":
            assert item["value"] in pdf_text, item
    assert workbook["PBC coverage"].max_row == 12
    assert workbook["Assurance criteria"].max_row == 59
    assert workbook["Metadata"].max_row == 19
    assert all(not workbook["Reviewer signoff"].cell(row, 2).value for row in range(3, 8))
    with pdfplumber.open(pdf_path) as pdf:
        for number, page in enumerate(pdf.pages, 1):
            bad = [char for char in page.chars if char["x0"] < -1 or char["x1"] > page.width + 1 or char["top"] < -1 or char["bottom"] > page.height + 1]
            assert not bad, f"Text outside page {number}: {bad[:3]}"
    if name == "populated":
        assert "SYNTHETIC DEMO" in pdf_text
        assert workbook["Summary"]["B5"].value.startswith("SYNTHETIC DEMO")
        assert isinstance(workbook["Accounts"]["G3"].value, (int, float))
        assert isinstance(workbook["Accounts"]["H3"].value, (int, float))
    print(f"{name}: {len(reader.pages)} PDF pages; all PBC/criteria/metadata present; XLSX and bounds verified")


if __name__ == "__main__":
    verify("populated")
    verify("sparse-long")
